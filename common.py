"""Общие функции: конфиг, фильтр по ключевым словам, сохранение в CSV/Excel без дублей."""
import csv
import hashlib
import re
import sqlite3
from datetime import datetime
from pathlib import Path

import yaml

BASE = Path(__file__).resolve().parent
DATA = BASE / "data"
DATA.mkdir(exist_ok=True)

HEADERS = {
    "User-Agent": "Mozilla/5.0 (X11; Ubuntu; Linux x86_64; rv:130.0) Gecko/20100101 Firefox/130.0",
    "Accept-Language": "ru-RU,ru;q=0.9,en;q=0.5",
}

FIELDS = ["found_at", "source", "title", "company", "city", "phone", "email",
          "website", "url", "text", "score", "notes"]


def _merge(defaults, user):
    """Недостающие ключи из defaults; значения пользователя (и его списки) не трогаем."""
    for k, v in defaults.items():
        if k not in user:
            user[k] = v
        elif isinstance(v, dict) and isinstance(user[k], dict):
            _merge(v, user[k])
    return user


def load_config():
    if not (BASE / "config.yaml").exists() and (BASE / "config.example.yaml").exists():
        import shutil
        shutil.copy(BASE / "config.example.yaml", BASE / "config.yaml")  # первый запуск
    with open(BASE / "config.yaml", encoding="utf-8") as f:
        cfg = yaml.safe_load(f) or {}
    example = BASE / "config.example.yaml"
    if example.exists():
        with open(example, encoding="utf-8") as f:
            cfg = _merge(yaml.safe_load(f) or {}, cfg)
    return cfg


def matches(text, cfg):
    """True, если есть ключевое слово и нет стоп-слова."""
    t = (text or "").lower()
    if any(s.lower() in t for s in cfg.get("stop_words") or []):
        return False
    return any(k.lower() in t for k in cfg.get("keywords") or [])


EMAIL_RE = re.compile(r"[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,}")
PHONE_RE = re.compile(r"(?:\+7|8)[\s\-(]*\d{3}[\s\-)]*\d{3}[\s\-]*\d{2}[\s\-]*\d{2}")
BAD_EMAIL_END = (".png", ".jpg", ".jpeg", ".gif", ".svg", ".webp", ".js", ".css")


def find_emails(text):
    out = []
    for e in EMAIL_RE.findall(text or ""):
        e = e.strip(".").lower()
        if not e.endswith(BAD_EMAIL_END) and "example" not in e and e not in out:
            out.append(e)
    return out


def norm_phone(p):
    d = re.sub(r"\D", "", p or "")
    if len(d) == 11 and d[0] in "78":
        return "+7" + d[1:]
    if len(d) == 10:
        return "+7" + d
    return p.strip() if p else ""


def find_phones(text):
    out = []
    for p in PHONE_RE.findall(text or ""):
        n = norm_phone(p)
        if n not in out:
            out.append(n)
    return out


class Store:
    """Пишет лиды в data/<name>.csv, пропуская уже виденные (дедуп в data/seen.db)."""

    def __init__(self, name):
        self.path = DATA / f"{name}.csv"
        self.db = sqlite3.connect(DATA / "seen.db")
        self.db.execute("CREATE TABLE IF NOT EXISTS seen (key TEXT PRIMARY KEY, at TEXT)")
        new_file = not self.path.exists() or self.path.stat().st_size == 0
        self.f = open(self.path, "a", newline="", encoding="utf-8-sig")
        self.w = csv.DictWriter(self.f, fieldnames=FIELDS, extrasaction="ignore")
        if new_file:
            self.w.writeheader()
        self.added = 0

    def add(self, lead, key=None):
        key = key or lead.get("url") or lead.get("website") or (lead.get("company", "") + lead.get("phone", ""))
        h = hashlib.sha1(f"{self.path.stem}|{key}".encode()).hexdigest()
        if self.db.execute("SELECT 1 FROM seen WHERE key=?", (h,)).fetchone():
            return False
        self.db.execute("INSERT INTO seen VALUES (?, ?)", (h, datetime.now().isoformat()))
        lead.setdefault("found_at", datetime.now().strftime("%Y-%m-%d %H:%M"))
        if lead.get("text"):
            lead["text"] = re.sub(r"\s+", " ", lead["text"])[:1000]
        self.w.writerow(lead)
        self.added += 1
        return True

    def close(self):
        self.db.commit()
        self.db.close()
        self.f.close()
        print(f"  -> новых лидов: {self.added}, файл: {self.path}")


def export_excel():
    """Собирает все CSV из data/ в один leads.xlsx (по листу на источник)."""
    from openpyxl import Workbook
    from openpyxl.styles import Font
    wb = Workbook()
    wb.remove(wb.active)
    files = sorted(DATA.glob("*.csv"))
    for p in files:
        ws = wb.create_sheet(p.stem[:31])
        with open(p, encoding="utf-8-sig") as f:
            for i, row in enumerate(csv.reader(f)):
                ws.append(row)
                if i == 0:
                    for c in ws[1]:
                        c.font = Font(bold=True)
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
    if not files:
        print("Нет данных для экспорта")
        return
    out = DATA / "leads.xlsx"
    wb.save(out)
    print(f"Excel сохранён: {out}")
